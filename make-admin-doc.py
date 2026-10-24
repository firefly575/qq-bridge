# -*- coding: utf-8 -*-
"""生成《QQ机器人-管理员指令.docx》（默认输出到本脚本同目录），内容以当前代码实际支持为准。

用法：python make-admin-doc.py [--out 路径]
依赖：pip install python-docx
"""
import io
import sys
from datetime import datetime

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

import os as _os
_HERE = _os.path.dirname(_os.path.abspath(__file__))
# 默认输出到脚本同目录，避免依赖某个用户的桌面路径；可用 --out 覆盖
OUT = _os.path.join(_HERE, "QQ机器人-管理员指令.docx")
if "--out" in sys.argv:
    OUT = sys.argv[sys.argv.index("--out") + 1]

doc = Document()

# ---------- 样式 ----------
style = doc.styles["Normal"]
style.font.name = "微软雅黑"
style.font.size = Pt(10.5)
try:
    style.element.rPr.rFonts.set(
        __import__("docx").oxml.ns.qn("w:eastAsia"), "微软雅黑")
except Exception:
    pass

for name in ("Heading 1", "Heading 2", "Title"):
    try:
        doc.styles[name].font.name = "微软雅黑"
    except Exception:
        pass


def h1(text):
    p = doc.add_heading(text, level=1)
    return p


def para(text, bold=False, size=None, color=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)
    return p


def bullet(text):
    return doc.add_paragraph(text, style="List Bullet")


def code(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9.5)
    p.paragraph_format.space_after = Pt(2)
    return p


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, htext in enumerate(headers):
        hdr[i].text = ""
        run = hdr[i].paragraphs[0].add_run(htext)
        run.bold = True
        run.font.size = Pt(10)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(val))
            run.font.size = Pt(9.5)
    if widths:
        for r in t.rows:
            for i, w in enumerate(widths):
                r.cells[i].width = w
    # 表格后留一点空隙，避免下一段贴着表格
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)
    return t


# ---------- 封面信息 ----------
title = doc.add_heading("QQ 机器人 · 管理员指令", level=0)
para("qq-bridge 硬执行命令参考（所有指令只有 ownerQQ 可用）", bold=True)
para(f"管理员 QQ：<管理员QQ>    模式：reserved2    更新：{datetime.now():%Y-%m-%d %H:%M}")
doc.add_paragraph()

# ---------- 使用前提 ----------
h1("使用前提")
bullet("这些指令由 qq-bridge 直接执行，不经过 AI 模型、不消耗 token，也不会被群友诱导。")
bullet("只有管理员 QQ（ownerQQ <管理员QQ>）发才算数；其他人发同样内容会被忽略或拦截。")
bullet("指令必须以 / 开头、整条消息只有指令本身（前后不要带别的字）。")
bullet("不认识的 /xxx 会原样转给 DSH，那就变成模型回答了（会花 token）。")

# ---------- 指令总表 ----------
h1("指令一览")

table(
    ["指令", "参数", "作用"],
    [
        ["/reset", "（别名 /new）", "重置当前会话：清空上下文重新开始"],
        ["/status", "—", "查看状态：会话、白名单、角色、当前模型"],
        ["/role <名字>", "或 /role off", "切换人格；角色文件来自 qq-bridge/roles/"],
        ["/silent", "（别名 /quiet）", "静默模式：群友消息不回复，仅管理员可对话"],
        ["/active", "（别名 /speak）", "退出静默，恢复正常回复"],
        ["/model", "见下节", "在 API 模型与本地 Ollama 之间切换"],
        ["/cover", "见下节", "调整自动翻唱的处理时长与音色"],
        ["/sticker", "见下节", "调整自动发表情的概率与开关"],
        ["/draw", "见下节", "切换生图模型（qwen / seedream / seedream-pro / sdcpp）、分辨率、每日出图上限"],
        ["/video", "见下节", "视频生成：一次一片、480P、2~30 秒、每日上限、测试片"],
    ],
    widths=[Pt(90), Pt(110), Pt(270)],
)

# ---------- /model ----------
h1("/model —— 模型切换")

table(
    ["指令", "作用"],
    [
        ["/model", "查看当前用的哪个模型"],
        ["/model api", "切回 API（DSH 会话，工具齐全，消耗 token）"],
        ["/model local", "切到本地 Ollama（0 token，但没有工具）"],
        ["/model list", "列出本机 Ollama 已装模型"],
        ["/model 1.5b", "按名字切到本地模型（支持简写，自动匹配全名）"],
    ],
    widths=[Pt(110), Pt(360)],
)

para("说明：")
bullet("本地模式不走 DSH 会话，桥接直接把消息发给本机 Ollama 的 /api/chat，再代为发送回复。")
bullet("切换立即生效，不需要重启；当前选择保存在 state/model-mode.json。")
bullet("切之前桥接会先发一个超短请求「试加载」，成功才真正切过去。")
bullet("切到本地后工具全没了：画图、翻唱、语音、表情都不可用。要用这些能力请 /model api。")

# ---------- /cover ----------
h1("/cover —— 自动翻唱设置")

table(
    ["指令", "作用"],
    [
        ["/cover", "查看当前处理时长、音色、可用音色列表"],
        ["/cover 30", "设置默认翻唱时长（10~180 秒，自动夹紧）"],
        ["/cover voice", "列出所有音色预设 + 当前用的是哪个"],
        ["/cover voice v2", "切换到指定音色（v0 / v1 / v2 …）"],
    ],
    widths=[Pt(110), Pt(360)],
)

para("默认值：15 秒 / 音色 v2 / 伴奏音量 2.16（人声 1.0）", bold=True)
bullet("音色预设放在 qq-bridge/audio/voices/，文件名就是预设名。")
bullet("换音色会同时影响「翻唱」和「说话」两条链路。")
bullet("设置存在 state/cover-settings.json，重启不丢。")

# ---------- /sticker ----------
h1("/sticker —— 自动发表情设置")

table(
    ["指令", "作用"],
    [
        ["/sticker", "查看开关状态、当前概率、收藏设置"],
        ["/sticker 35", "把自动发表情概率设为 35%（写 35% 或 0.35 都行）"],
        ["/sticker off", "关闭自动发表情"],
        ["/sticker on", "重新开启（沿用上次概率）"],
        ["/sticker collect off", "关闭「别人发表情自动收藏」"],
        ["/sticker collect on", "重新开启自动收藏"],
    ],
    widths=[Pt(120), Pt(350)],
)

para("说明：")
bullet("概率越高越活跃；0.22 约等于每 4~5 次回复带一个表情。")
bullet("只有和当前话题匹配（备注/标签/收藏语境有重合）才会发，不是纯随机。")
bullet("设置存在 state/sticker-settings.json，重启不丢。")

# ---------- /draw ----------
h1("/draw —— 生图设置（模型 / 分辨率 / 每日上限）")

table(
    ["指令", "作用"],
    [
        ["/draw", "查看当前模型、型号、分辨率、今日已用几张、还能画几张"],
        ["/draw model qwen", "切成百炼 Qwen-Image（与语音共用同一把阿里云百炼 key）"],
        ["/draw model seedream", "切成火山 Seedream 5.0 Flash（快档，一张约 10 秒）"],
        ["/draw model seedream-pro", "切成火山 Seedream 5.0 Pro（高端档，一张 18~40 秒，中文书法/印章更准）"],
        ["/draw model sdcpp", "切成本地 sd-cpp（纯 CPU，慢，不联网不花钱）"],
        ["/draw res 1080", "分辨率切到 1080P（1920×1080）"],
        ["/draw res 720", "分辨率切到 720P（1280×720，更快更便宜）"],
        ["/draw limit 10", "每日出图上限设为 10 张（0 = 当天不再生成新图）"],
    ],
    widths=[Pt(160), Pt(310)],
)

para("说明：")
bullet("分辨率只支持 720P / 1080P 两档；云端按 16:9 出图，本地 sd-cpp 仍是 512×512。")
bullet("每日上限默认 10 张，按自然日归零；缓存命中的重复提示词不占额度、不花钱。")
bullet("额度用完时不会发起请求，会直接告诉模型「今天的生图额度用完了」。")
bullet("设置存在 state/draw-settings.json，用量记录在 state/draw-usage.json，重启不丢。")
bullet("两个 seedream 档位各用一把火山方舟 key：Flash → state/ark-key.txt（或环境变量 ARK_API_KEY）；Pro → state/ark-key-2.txt（或 ARK_API_KEY_PRO）。百炼那把 key 与它们都不通用。")
bullet("实测耗时：Flash 720P 约 11 秒 / 1080P 约 9~10 秒；Pro 1080P 约 18~40 秒。方舟默认输出 jpeg（约 0.3~0.5MB），比 png 更省流量。")
bullet("模型型号写在 config.json 的 socialV2.draw.seedream / seedreamPro 里，要换版本改 model 即可。")
bullet("当前型号：qwen = qwen-image-3.0；Flash = doubao-seedream-5-0-flash-260915；Pro = doubao-seedream-5-0-pro-260628。")

# ---------- /video ----------
h1("/video —— 视频生成设置（一次一片｜480P｜2~30 秒）")

table(
    ["指令", "作用"],
    [
        ["/video", "查看开关、后端、型号、分辨率、时长、今日已出几条"],
        ["/video model wan", "阿里云百炼 wan3.0-video（**默认**）"],
        ["/video model wan-prime", "阿里云百炼 wan3.0-video-prime（同接口另一档）"],
        ["/video duration 4", "时长，**2~30 秒**（默认 4，30 秒也支持）"],
        ["/video limit 5", "每日条数上限（默认 5，0 = 当天不再生成）"],
        ["/video off", "关闭视频生成（模型不再拿到 qq_video 工具）"],
        ["/video on", "重新开启"],
        ["/video test", "测试片：生成 2 秒 → 本地剪成 **0.5 秒** 480P 再发（只一条）"],
    ],
    widths=[Pt(150), Pt(320)],
)

para("硬规则（重要）：")
bullet("**一次请求只生成一片**：失败就是失败，绝不自动换尺寸/换参数再提交第二个任务（重试=重复扣费）。")
bullet("默认 **480P（854×480）+ 4 秒**；时长接受 **2~30 秒**，越长越贵（按秒计费）。")
bullet("/video test 是唯一的短时长入口：**百炼接口最短只能 2 秒**（实测 duration 必须在 [2,30]，传 0.5 直接报错）。所以测试片是「生成 2 秒 → 用本地 ffmpeg 剪成 0.5 秒再发」，**只生成一条、不会重试**。")
bullet("seedance（火山方舟）**已搁置**：模型未开通，`/video model seedance` 会直接被拒绝，不会发任何请求。")
bullet("百炼按量计费，账号欠费时返回「余额不足」——**欠费会连图片一起拦**，不只视频。")
bullet("每日上限默认 5 条，按自然日归零；命中最缓存的同一提示词不占额度、不再扣费。")
bullet("在 QQ 里说「生成一段视频：……」即会触发；设置存 state/video-settings.json，用量在 state/video-usage.json。")
bullet("成品留在 videos/ 里当缓存（视频重生成很贵，所以不像翻唱那样发完即删）。")

# ---------- 不用指令的能力 ----------
h1("不需要指令的能力（自动触发）")

table(
    ["你做什么", "它会做什么"],
    [
        ["发音乐链接 + 「翻唱」", "自动下载 → 分离人声伴奏 → 换嗓 → 混音 → 发语音"],
        ["发音频文件 + 「翻唱」", "同上，直接用文件"],
        ["「翻唱 歌名」", "按歌名去 QQ 音乐找歌源，拿到音频再翻唱"],
        ["别人发表情", "自动收藏进表情库（去重 + 每小时限 8 个）"],
        ["聊天语境合适时", "按 /sticker 设定的概率，自己发一个对味的表情"],
    ],
    widths=[Pt(150), Pt(320)],
)

para("支持的链接平台：")
code("抖音 / 快手 / 汽水音乐 / QQ音乐 / 网易云 / B站 / YouTube / 微博 / 小红书 …")
para("有链接就按链接取；只给歌名时先去 QQ 音乐搜（优先免费版）。QQ 音乐全是 VIP 时转抖音搜，取视频里的音轨。")

# ---------- 语音能力 ----------
h1("语音与翻唱的技术现状")

table(
    ["能力", "实现", "耗时", "费用"],
    [
        ["说话", "百炼 qwen-audio-3.0-tts-flash + 复刻音色", "1~3 秒", "约 0.006 元/条"],
        ["翻唱", "demucs 分离 + Seed-VC 换嗓（本机 CPU）", "15 秒音频约 2~3 分钟", "免费"],
    ],
    widths=[Pt(70), Pt(260), Pt(110), Pt(90)],
)

para("语气：", bold=True)
bullet("说话时会按文字内容自动选语气（开心/难过/撒娇/认真等），普通陈述句保持自然。")
bullet("线索不够强就不加语气——硬凹会显得做作。")

para("翻唱规则：", bold=True)
bullet("同一首歌可以反复翻唱；同时只会跑一个，第二首会排队。")
bullet("生成的文件发完就删，不占磁盘。")

# ---------- 容易踩的点 ----------
h1("几个容易踩的点")
for text in [
    "「翻唱」和「唱歌」是同一件事：现在只走翻唱（换嗓），不再有「念歌词带调」。",
    "翻唱很吃 CPU：15 秒要 2~3 分钟，期间机器会有点卡。这是没显卡的代价。",
    "只给歌词（没链接没文件）不会唱：没有旋律来源，它会要你给首歌。",
    "换个音色要重新听效果：不同音色的中位音高不同，翻唱的音区会跟着变。",
    "表情自动收藏有上限（每小时 8 个），刷屏也不会把你收藏撑爆。",
    "指令打错会被当成普通消息发给模型（走 token）；指令表在上面，照着打。",
]:
    doc.add_paragraph(text, style="List Number")

# ---------- 文件位置 ----------
h1("文件与地址")
code("程序目录：D:\\dsh\\qq-bridge")
code("控制台：http://127.0.0.1:3100")
code("人格文件：qq-bridge\\roles\\<名字>.md")
code("音色预设：qq-bridge\\audio\\voices\\<名字>.wav")
code("配置文件：qq-bridge\\config.json")
code("运行状态：qq-bridge\\state\\（model-mode / cover-settings / sticker-settings / stickers）")
code("日志：qq-bridge\\state\\bridge.log、qq-activity.log")

doc.add_paragraph()
para("这份文档由脚本按当前代码生成，改了指令记得同步更新。", bold=False, size=9)

doc.save(OUT)
print("已生成:", OUT)
print("段落数:", len(doc.paragraphs), " 表格数:", len(doc.tables))
